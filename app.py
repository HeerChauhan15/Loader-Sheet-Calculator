import streamlit as st
import pandas as pd
import numpy as np
import os
import re
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

st.set_page_config(
    page_title="Insurance Premium Calculator",
    page_icon="💰",
    layout="wide"
)

st.title("💰 Aviva GCL Insurance Premium Calculator")
st.markdown("Select plan details below")

FILE_MAP = {
    ("Single Life", "Home Loan"): "Homeloan Single Life.xlsx",
    ("Single Life", "LAP"):       "LAP Single Life.xlsx",
    ("Joint Life",  "Home Loan"): "Homeloan Joint Life.xlsx",
    ("Joint Life",  "LAP"):       "LAP Joint Life.xlsx",
}

# Maximum age allowed for any borrower at the end of the loan tenure.
# For Joint Life, the loan tenure used for BOTH borrowers is capped so that
# neither borrower's age + tenure exceeds this limit (e.g. a 60-year-old
# Co Borrower limits the tenure to 5 years for both Main and Co Borrower).
MAX_AGE = 65

# GST is fixed and always applied on top of the Loader-adjusted rate.
GST_RATE_FIXED = 18.0

# ============================================
# LOADER + GST FORMULA (applied to every rate before premium is computed):
#   Rate After Loader = Base Rate / (1 - Loader% / 100)
#   Final Rate         = Rate After Loader x (1 + GST% / 100)
# ============================================
def apply_loader_and_gst(base_rate, loader_pct, gst_pct=GST_RATE_FIXED):
    after_loader = base_rate / (1 - (loader_pct / 100.0))
    final_rate = after_loader * (1 + (gst_pct / 100.0))
    return final_rate


def load_rate_table(life_type, loan_type):
    fname = FILE_MAP[(life_type, loan_type)]
    if not os.path.exists(fname):
        raise FileNotFoundError(
            f"File not found: '{fname}' — Please make sure this file is in the GitHub repo."
        )

    raw = pd.read_excel(fname, sheet_name="Sheet1", header=None)

    header_row = None
    for i, row in raw.iterrows():
        for val in row.values:
            if isinstance(val, str) and "AGE" in val.upper():
                header_row = i
                break
        if header_row is not None:
            break

    if header_row is None:
        raise ValueError("Could not find AGE/TERM header row in the file.")

    df = pd.read_excel(fname, sheet_name="Sheet1", header=header_row)
    df.columns = [str(c).strip() for c in df.columns]

    age_col = df.columns[0]
    df = df.dropna(subset=[age_col])
    df[age_col] = pd.to_numeric(df[age_col], errors='coerce')
    df = df.dropna(subset=[age_col])
    df[age_col] = df[age_col].astype(int)
    df = df.set_index(age_col)

    tenure_map = {}
    for col in df.columns:
        try:
            tenure_map[int(float(col))] = col
        except Exception:
            pass

    return df, tenure_map


def get_rate(df, tenure_map, age, tenure):
    if age not in df.index:
        raise ValueError(f"Age {age} not found in rate table.")
    if tenure not in tenure_map:
        raise ValueError(f"Tenure {tenure} yrs not found in rate table.")
    return float(df.loc[age, tenure_map[tenure]])


def find_column(df, target):
    """Exact (case/space-insensitive) match — used for Single Life."""
    target_norm = target.strip().lower().replace(" ", "")
    for col in df.columns:
        col_norm = str(col).strip().lower().replace(" ", "")
        if col_norm == target_norm:
            return col
    return None


def _normalize(s):
    return re.sub(r'[\s_\-]+', '', str(s).lower())


def detect_person(norm):
    """Decide whether a (normalized) column name belongs to Main Borrower or Co Borrower."""
    if 'mainborrower' in norm or norm.startswith('mb') or 'borrower1' in norm:
        return 'main'
    if 'coborrower' in norm or norm.startswith('cb') or 'borrower2' in norm:
        return 'co'
    # fallback: trailing 1 / 2 like Name1 / Name2, Age1 / Age2
    if norm.endswith('1'):
        return 'main'
    if norm.endswith('2'):
        return 'co'
    return None


def detect_field(norm):
    if 'name' in norm:
        return 'name'
    if 'tenure' in norm:
        return 'tenure'
    if 'age' in norm:
        return 'age'
    return None


def map_joint_columns(df):
    """
    Flexibly detect Main Borrower / Co Borrower Name/Age/Tenure columns
    regardless of exact header wording (Main Borrower Name, MB Name, Name1,
    Borrower 1 Name, etc.)
    """
    mapping = {}  # (person, field) -> actual column name
    for col in df.columns:
        norm = _normalize(col)
        field = detect_field(norm)
        if field is None:
            continue
        person = detect_person(norm)
        if person is None:
            continue
        key = (person, field)
        if key not in mapping:
            mapping[key] = col
    return mapping


# ============================================
# DROPDOWNS
# ============================================

col1, col2 = st.columns(2)
with col1:
    life_type = st.selectbox("Select Life Type", ["Single Life", "Joint Life"])
with col2:
    loan_type = st.selectbox("Select Loan Type", ["Home Loan", "LAP"])

# ============================================
# SHARED LOADER % — applied to every rate (Manual Single, Manual Joint, Bulk)
# before computing premium. GST @ 18% is then applied automatically on top.
# ============================================
st.subheader("⚙️ Loader Setting")
loader_pct = st.number_input(
    "Loader % (applied to all rate lookups below; GST @ 18% is then added automatically)",
    min_value=0.0,
    max_value=99.99,
    value=None,
    step=1.0,
    placeholder="Enter loader %",
    key="shared_loader_pct"
)
st.info(f"ℹ️ GST @ {GST_RATE_FIXED}% will be added automatically after the Loader — it is not editable.")

loader_ready = loader_pct is not None and loader_pct < 100

# ============================================
# SUM ASSURED RANGE — rates in the backend files are per ₹1,00,000
# (The actual Sum Assured input widgets are now defined separately inside
# the Manual and Bulk sections below, so changing one does not affect
# the others.)
# ============================================
if loan_type == "Home Loan":
    sa_min, sa_max = 100000, 6000000
else:
    sa_min, sa_max = 100000, 4000000

st.divider()

# ============================================
# MANUAL SECTION
# ============================================

st.subheader("🔢 Manual Rate Lookup")

if not loader_ready:
    st.warning("⚠ Set a valid Loader % (below 100) above to unlock rate lookup.")

if loan_type == "Home Loan":
    min_tenure, max_tenure = 5, 25
else:
    min_tenure, max_tenure = 2, 10

if life_type == "Single Life":
    col3, col4 = st.columns(2)
    with col3:
        age = st.number_input("Enter Age", min_value=18, max_value=65, value=30, step=1)
    with col4:
        tenure = st.number_input(
            "Enter Tenure",
            min_value=min_tenure,
            max_value=max_tenure,
            value=min_tenure,
            step=1
        )
        st.caption("📅 Tenure is in Years")

    sum_assured_manual_single = st.number_input(
        "Select Sum Assured (₹)",
        min_value=sa_min,
        max_value=sa_max,
        value=sa_min,
        step=100000,
        help=f"For {loan_type}, Sum Assured must be between ₹{sa_min:,} and ₹{sa_max:,}.",
        key="sa_manual_single"
    )

    if st.button("Get Rate", type="primary", disabled=not loader_ready):
        try:
            df_rates, tenure_map = load_rate_table(life_type, loan_type)
            base_rate = get_rate(df_rates, tenure_map, age, tenure)
            final_rate = apply_loader_and_gst(base_rate, loader_pct)
            premium = final_rate * (sum_assured_manual_single / 100000)
            st.success(
                f"✅ {life_type} | {loan_type} | Age {age} | Tenure {tenure} yrs | "
                f"Sum Assured ₹{sum_assured_manual_single:,} | Loader {loader_pct}% | GST {GST_RATE_FIXED}%"
            )
            col_a, col_b, col_c = st.columns(3)
            with col_a:
                st.metric("Base Rate (per ₹1,00,000)", f"₹ {base_rate:,.2f}")
            with col_b:
                st.metric("Rate after Loader + GST", f"₹ {final_rate:,.2f}")
            with col_c:
                st.metric("Premium (for selected Sum Assured)", f"₹ {premium:,.2f}")
        except Exception as e:
            st.error(f"Error: {e}")

else:
    st.markdown("**Main Borrower**")
    mcol1, mcol2 = st.columns(2)
    with mcol1:
        main_age = st.number_input("Age", min_value=18, max_value=65, value=30, step=1, key="main_age_manual")
    with mcol2:
        main_tenure = st.number_input(
            "Tenure", min_value=min_tenure, max_value=max_tenure,
            value=min_tenure, step=1, key="main_tenure_manual"
        )
    st.caption("📅 Tenure is in Years")

    st.markdown("**Co Borrower**")
    ccol1, ccol2 = st.columns(2)
    with ccol1:
        co_age = st.number_input("Age", min_value=18, max_value=65, value=30, step=1, key="co_age_manual")
    with ccol2:
        co_tenure = st.number_input(
            "Tenure", min_value=min_tenure, max_value=max_tenure,
            value=min_tenure, step=1, key="co_tenure_manual"
        )
    st.caption(f"📅 Tenure is in Years. Maximum age allowed at end of tenure is {MAX_AGE} — the lower of the two borrowers' allowed tenures is used for both.")

    sum_assured_manual_joint = st.number_input(
        "Select Sum Assured (₹)",
        min_value=sa_min,
        max_value=sa_max,
        value=sa_min,
        step=100000,
        help=f"For {loan_type}, Sum Assured must be between ₹{sa_min:,} and ₹{sa_max:,}.",
        key="sa_manual_joint"
    )

    if st.button("Get Rate", type="primary", disabled=not loader_ready):
        try:
            df_rates, tenure_map = load_rate_table(life_type, loan_type)

            # Loan tenure is shared — cap it so neither borrower's age + tenure
            # exceeds MAX_AGE, then use the same (lower) tenure for both.
            main_age_cap = MAX_AGE - main_age
            co_age_cap = MAX_AGE - co_age
            effective_tenure = min(main_tenure, co_tenure, main_age_cap, co_age_cap)

            if effective_tenure < min_tenure:
                raise ValueError(
                    f"Effective tenure ({effective_tenure} yrs) falls below the minimum "
                    f"allowed tenure ({min_tenure} yrs) because of the age limit (max age {MAX_AGE})."
                )

            rate_main_base = get_rate(df_rates, tenure_map, main_age, effective_tenure)
            rate_main_final = apply_loader_and_gst(rate_main_base, loader_pct)
            premium_main = rate_main_final * (sum_assured_manual_joint / 100000)

            rate_co_base = get_rate(df_rates, tenure_map, co_age, effective_tenure)
            rate_co_final = apply_loader_and_gst(rate_co_base, loader_pct)
            premium_co = rate_co_final * (sum_assured_manual_joint / 100000)

            total_premium = premium_main + premium_co

            if effective_tenure < max(main_tenure, co_tenure):
                st.info(
                    f"ℹ️ Tenure capped to {effective_tenure} yrs for both borrowers because of the "
                    f"age limit (max age {MAX_AGE}). Main Borrower entered {main_tenure} yrs, "
                    f"Co Borrower entered {co_tenure} yrs."
                )

            st.success(
                f"✅ {life_type} | {loan_type} | Sum Assured ₹{sum_assured_manual_joint:,} | "
                f"Loader {loader_pct}% | GST {GST_RATE_FIXED}% | "
                f"Main Borrower: Age {main_age}, Tenure used {effective_tenure} yrs | "
                f"Co Borrower: Age {co_age}, Tenure used {effective_tenure} yrs"
            )

            col_a, col_b, col_c = st.columns(3)
            with col_a:
                st.metric("Main Borrower Premium", f"₹ {premium_main:,.2f}")
            with col_b:
                st.metric("Co Borrower Premium", f"₹ {premium_co:,.2f}")
            with col_c:
                st.metric("Total Premium", f"₹ {total_premium:,.2f}")
        except Exception as e:
            st.error(f"Error: {e}")

st.divider()

# ============================================
# EXCEL UPLOAD SECTION
# ============================================

st.subheader("📂 Upload Member Data for Bulk Rate Lookup")

if life_type == "Single Life":
    st.markdown(
        "Your Excel must have at least: **Name**, **Age**, **Tenure** (in years)."
    )
else:
    st.markdown(
        "Your Excel must have at least: **Main Borrower** (Name, Age, Tenure in years) and "
        "**Co Borrower** (Name, Age, Tenure in years). "
        f"Loan tenure is shared between borrowers — if either borrower's age + tenure would "
        f"exceed {MAX_AGE} years, the tenure is automatically capped for both borrowers."
    )

st.caption("Rates/premium shown are as per ₹1,00,000 Sum Assured.")

sum_assured_bulk = 100000

st.warning("⚠️ Please make sure you have selected **Life Type**, **Loan Type**, and a valid **Loader %** above before uploading your Excel file.")

uploaded_file = st.file_uploader("Upload Excel File", type=["xlsx"], disabled=not loader_ready)

if uploaded_file is not None and loader_ready:
    try:
        df = pd.read_excel(uploaded_file)
        df.columns = [str(c).strip() for c in df.columns]

        st.subheader("Uploaded Data Preview")
        st.dataframe(df.head())

        if loan_type == "Home Loan":
            min_t, max_t = 5, 25
        else:
            min_t, max_t = 2, 10

        df_rates, tenure_map = load_rate_table(life_type, loan_type)

        # ============================================
        # SINGLE LIFE
        # ============================================
        if life_type == "Single Life":
            name_col = find_column(df, "Name")
            age_col = find_column(df, "Age")
            tenure_col = find_column(df, "Tenure")

            if not name_col or not age_col or not tenure_col:
                raise ValueError("Excel must contain mandatory columns: Name, Age, Tenure")

            df[age_col] = pd.to_numeric(df[age_col], errors='coerce')
            df[tenure_col] = pd.to_numeric(df[tenure_col], errors='coerce')

            if df[tenure_col].dropna().median() > 30:
                st.info("ℹ️ Tenure values look like months — auto-converting to years.")
                df[tenure_col] = (df[tenure_col] / 12).round(0).astype('Int64')
            else:
                df[tenure_col] = df[tenure_col].round(0).astype('Int64')

            df[age_col] = df[age_col].round(0).astype('Int64')
            df[tenure_col] = df[tenure_col].clip(lower=min_t, upper=max_t)

            premiums = []
            statuses = []
            for idx, row in df.iterrows():
                try:
                    r_age = int(row[age_col])
                    r_tenure = int(row[tenure_col])
                    r_base = get_rate(df_rates, tenure_map, r_age, r_tenure)
                    r_final = apply_loader_and_gst(r_base, loader_pct)
                    premium = round(r_final * (sum_assured_bulk / 100000), 2)
                    premiums.append(premium)
                    statuses.append("✅")
                except Exception as e:
                    premiums.append(None)
                    statuses.append(f"❌ {e}")

            df["Premium"] = premiums
            df["Status"] = statuses

            core_cols = [name_col, age_col, tenure_col, "Premium"]
            extra_cols = [c for c in df.columns if c not in core_cols]
            df = df[core_cols + extra_cols]

            total_premium = pd.to_numeric(pd.Series(premiums), errors='coerce').sum()
            st.metric("💰 Grand Total Premium", f"₹ {total_premium:,.2f}")

            st.subheader("Rate Lookup Output")
            st.dataframe(df, use_container_width=True)

            total_row = {c: "" for c in df.columns}
            total_row[name_col] = "TOTAL PREMIUM"
            total_row["Premium"] = round(total_premium, 2)
            df_out = pd.concat([df, pd.DataFrame([total_row])], ignore_index=True)

            output_file = "Rate_Output.xlsx"
            df_out.to_excel(output_file, index=False)

            with open(output_file, "rb") as file:
                st.download_button(
                    label="⬇ Download Output Excel",
                    data=file,
                    file_name=output_file,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

        # ============================================
        # JOINT LIFE
        # ============================================
        else:
            mapping = map_joint_columns(df)

            required_keys = [
                ('main', 'name'), ('main', 'age'), ('main', 'tenure'),
                ('co', 'name'), ('co', 'age'), ('co', 'tenure'),
            ]
            missing = [f"{p.capitalize()} Borrower {f.capitalize()}" for (p, f) in required_keys if (p, f) not in mapping]

            if missing:
                raise ValueError(
                    "Could not detect these mandatory columns: " + ", ".join(missing) +
                    ". Please make sure your Excel has Main Borrower & Co Borrower "
                    "Name/Age/Tenure columns (any reasonable naming works, e.g. "
                    "'Main Borrower Age', 'MB Age', 'Age1')."
                )

            main_name_col = mapping[('main', 'name')]
            main_age_col = mapping[('main', 'age')]
            main_tenure_col = mapping[('main', 'tenure')]
            co_name_col = mapping[('co', 'name')]
            co_age_col = mapping[('co', 'age')]
            co_tenure_col = mapping[('co', 'tenure')]

            for c in [main_age_col, main_tenure_col, co_age_col, co_tenure_col]:
                df[c] = pd.to_numeric(df[c], errors='coerce')

            for tcol, label in [(main_tenure_col, "Main Borrower Tenure"), (co_tenure_col, "Co Borrower Tenure")]:
                if df[tcol].dropna().median() > 30:
                    st.info(f"ℹ️ {label} values look like months — auto-converting to years.")
                    df[tcol] = (df[tcol] / 12).round(0).astype('Int64')
                else:
                    df[tcol] = df[tcol].round(0).astype('Int64')

            df[main_age_col] = df[main_age_col].round(0).astype('Int64')
            df[co_age_col] = df[co_age_col].round(0).astype('Int64')
            df[main_tenure_col] = df[main_tenure_col].clip(lower=min_t, upper=max_t)
            df[co_tenure_col] = df[co_tenure_col].clip(lower=min_t, upper=max_t)

            st.info(
                f"ℹ️ Loan tenure is shared between borrowers — if either borrower's age + tenure "
                f"would exceed {MAX_AGE} years, the tenure is automatically capped for both borrowers."
            )

            premium_main_list = []
            premium_co_list = []
            total_list = []
            tenure_used_list = []
            statuses = []

            for idx, row in df.iterrows():
                row_status = "✅"
                p_main = None
                p_co = None
                eff_tenure = None
                try:
                    m_age = int(row[main_age_col])
                    m_tenure = int(row[main_tenure_col])
                    c_age = int(row[co_age_col])
                    c_tenure = int(row[co_tenure_col])

                    # Loan tenure is shared — cap it so neither borrower's age + tenure
                    # exceeds MAX_AGE, then use the same (lower) tenure for both.
                    main_age_cap = MAX_AGE - m_age
                    co_age_cap = MAX_AGE - c_age
                    eff_tenure = min(m_tenure, c_tenure, main_age_cap, co_age_cap)

                    if eff_tenure < min_t:
                        raise ValueError(
                            f"Effective tenure ({eff_tenure} yrs) below minimum allowed "
                            f"({min_t} yrs) due to age limit (max age {MAX_AGE})"
                        )

                    rate_main_base = get_rate(df_rates, tenure_map, m_age, eff_tenure)
                    rate_main_final = apply_loader_and_gst(rate_main_base, loader_pct)
                    p_main = round(rate_main_final * (sum_assured_bulk / 100000), 2)

                    rate_co_base = get_rate(df_rates, tenure_map, c_age, eff_tenure)
                    rate_co_final = apply_loader_and_gst(rate_co_base, loader_pct)
                    p_co = round(rate_co_final * (sum_assured_bulk / 100000), 2)

                    if eff_tenure < max(m_tenure, c_tenure):
                        row_status = f"✅ (tenure capped to {eff_tenure} yrs due to age limit)"
                except Exception as e:
                    row_status = f"❌ {e}"

                premium_main_list.append(p_main)
                premium_co_list.append(p_co)
                total_list.append(round(p_main + p_co, 2) if (p_main is not None and p_co is not None) else None)
                tenure_used_list.append(eff_tenure)
                statuses.append(row_status)

            df["Main Borrower Premium"] = premium_main_list
            df["Co Borrower Premium"] = premium_co_list
            df["Tenure Used"] = tenure_used_list
            df["Total Premium"] = total_list
            df["Status"] = statuses

            core_cols = [
                main_name_col, main_age_col, main_tenure_col, "Main Borrower Premium",
                co_name_col, co_age_col, co_tenure_col, "Co Borrower Premium",
                "Tenure Used", "Total Premium"
            ]
            extra_cols = [c for c in df.columns if c not in core_cols]
            df_display = df[core_cols + extra_cols]

            grand_total = pd.to_numeric(pd.Series(total_list), errors='coerce').sum()
            st.metric("💰 Grand Total Premium", f"₹ {grand_total:,.2f}")

            st.subheader("Rate Lookup Output")
            st.dataframe(df_display, use_container_width=True)

            # ---- Build output with TOTAL PREMIUM row at the bottom ----
            total_row = {c: "" for c in df_display.columns}
            total_row[main_name_col] = "TOTAL PREMIUM"
            total_row["Total Premium"] = round(grand_total, 2)
            df_out = pd.concat([df_display, pd.DataFrame([total_row])], ignore_index=True)

            output_file = "Rate_Output.xlsx"

            # Write data starting from row 3 (1-indexed), leaving rows 1-2 for the 2-row header
            df_out.to_excel(output_file, index=False, header=False, startrow=2, sheet_name="Sheet1")

            wb = load_workbook(output_file)
            ws = wb["Sheet1"]

            bold = Font(bold=True)
            center = Alignment(horizontal="center", vertical="center")

            n_extra = len(extra_cols)
            # Column positions (1-indexed): Main group = cols 1-4 (Name, Age, Tenure, Premium),
            # Co group = cols 5-8 (Name, Age, Tenure, Premium),
            # Tenure Used = col 9, Total Premium = col 10, extras start at col 11
            main_start, main_end = 1, 4
            co_start, co_end = 5, 8
            tenure_used_col = 9
            total_col = 10
            extra_start = 11

            # Row 2: sub-headers (actual field names)
            row2_labels = ["Name", "Age", "Tenure", "Premium", "Name", "Age", "Tenure", "Premium",
                           "Tenure Used", "Total Premium"] + extra_cols
            for idx, label in enumerate(row2_labels, start=1):
                cell = ws.cell(row=2, column=idx, value=label)
                cell.font = bold
                cell.alignment = center

            # Row 1: group headers (merged)
            ws.merge_cells(start_row=1, start_column=main_start, end_row=1, end_column=main_end)
            ws.cell(row=1, column=main_start, value="MAIN BORROWER").font = bold
            ws.cell(row=1, column=main_start).alignment = center

            ws.merge_cells(start_row=1, start_column=co_start, end_row=1, end_column=co_end)
            ws.cell(row=1, column=co_start, value="CO BORROWER").font = bold
            ws.cell(row=1, column=co_start).alignment = center

            ws.merge_cells(start_row=1, start_column=tenure_used_col, end_row=2, end_column=tenure_used_col)
            ws.cell(row=1, column=tenure_used_col, value="TENURE USED").font = bold
            ws.cell(row=1, column=tenure_used_col).alignment = center
            ws.cell(row=2, column=tenure_used_col, value=None)

            ws.merge_cells(start_row=1, start_column=total_col, end_row=2, end_column=total_col)
            ws.cell(row=1, column=total_col, value="TOTAL PREMIUM").font = bold
            ws.cell(row=1, column=total_col).alignment = center
            # remove duplicate row2 label under merged Total Premium cell
            ws.cell(row=2, column=total_col, value=None)

            if n_extra:
                ws.merge_cells(start_row=1, start_column=extra_start, end_row=1, end_column=extra_start + n_extra - 1)
                ws.cell(row=1, column=extra_start, value="ADDITIONAL INFO").font = bold
                ws.cell(row=1, column=extra_start).alignment = center

            # Auto width
            for col_idx in range(1, total_col + n_extra + 1):
                col_letter = get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = 18

            wb.save(output_file)

            with open(output_file, "rb") as file:
                st.download_button(
                    label="⬇ Download Output Excel",
                    data=file,
                    file_name=output_file,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

    except Exception as e:
        st.error(f"Error: {e}")
